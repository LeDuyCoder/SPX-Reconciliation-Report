"""Excel workbook validation and streaming Main Data row conversion."""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Callable, Iterator

from .columns import COLUMNS


class ImportErrorMessage(Exception):
    """An actionable workbook or validation error for the UI."""


@dataclass
class ImportPreview:
    path: Path
    sheet_name: str
    row_count: int
    original_headers: list[str]
    missing_columns: list[str]
    unexpected_columns: list[str]


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\r", " ").replace("\n", " ")).strip().casefold()


def _open_sheet(path: Path):
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportErrorMessage("Thiếu thư viện openpyxl. Cài dependency của ứng dụng rồi thử lại.") from exc
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Main Data" not in book.sheetnames:
            names = ", ".join(book.sheetnames) or "(không có sheet)"
            book.close()
            raise ImportErrorMessage(f'Không tìm thấy sheet "Main Data". Các sheet hiện có: {names}')
        return book, book["Main Data"], "xlsx"
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ImportErrorMessage("Thiếu thư viện xlrd để đọc file .xls.") from exc
        try:
            book = xlrd.open_workbook(path, on_demand=True)
        except Exception as exc:
            raise ImportErrorMessage(f"Không mở được file Excel: {exc}") from exc
        if "Main Data" not in book.sheet_names():
            names = ", ".join(book.sheet_names()) or "(không có sheet)"
            book.release_resources()
            raise ImportErrorMessage(f'Không tìm thấy sheet "Main Data". Các sheet hiện có: {names}')
        return book, book.sheet_by_name("Main Data"), "xls"
    raise ImportErrorMessage("Định dạng không hỗ trợ. Hãy chọn file .xlsx hoặc .xls.")


def _rows(sheet, kind: str) -> Iterator[list[Any]]:
    if kind == "xlsx":
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    else:
        for index in range(sheet.nrows):
            yield sheet.row_values(index)


def _headers_first_nonempty(sheet, kind: str) -> tuple[list[Any], dict[str, int], int]:
    iterator = _rows(sheet, kind)
    for row_index, row in enumerate(iterator):
        if any(value is not None and str(value).strip() for value in row):
            original = ["" if value is None else str(value) for value in row]
            normalized = [normalize_header(value) for value in row]
            seen = set()
            for index, key in enumerate(normalized):
                if key and key in seen:
                    raise ImportErrorMessage(f'Header bị trùng sau chuẩn hóa: "{original[index]}".')
                if key:
                    seen.add(key)
            expected = {normalize_header(column.label): column.field for column in COLUMNS}
            mapping: dict[str, int] = {}
            for index, key in enumerate(normalized):
                if key in expected:
                    if expected[key] in mapping:
                        raise ImportErrorMessage(f'Header bị trùng sau chuẩn hóa: "{original[index]}".')
                    mapping[expected[key]] = index
            if not mapping:
                raise ImportErrorMessage("Không tìm thấy cột Main Data nào được hỗ trợ trong dòng tiêu đề.")
            return original, mapping, row_index
    raise ImportErrorMessage('Sheet "Main Data" không có dòng tiêu đề hoặc đang trống.')


def _headers(sheet, kind: str) -> tuple[list[Any], dict[str, int], int]:
    """Find the row that best matches known Main Data column names."""
    expected = {normalize_header(column.label): column.field for column in COLUMNS}
    best: tuple[list[Any], int, int] | None = None
    for row_index, row in enumerate(_rows(sheet, kind)):
        normalized = [normalize_header(value) for value in row]
        score = sum(1 for key in normalized if key in expected)
        if score and (best is None or score > best[2]):
            original = ["" if value is None else str(value) for value in row]
            best = original, row_index, score

    if best is None:
        raise ImportErrorMessage('Sheet "Main Data" has no recognized header row or is empty.')

    original, header_row, _ = best
    normalized = [normalize_header(value) for value in original]
    mapping: dict[str, int] = {}
    seen: set[str] = set()
    for index, key in enumerate(normalized):
        if not key:
            continue
        if key in seen:
            raise ImportErrorMessage(f'Duplicate header after normalization: "{original[index]}".')
        seen.add(key)
        if key in expected:
            mapping[expected[key]] = index
    return original, mapping, header_row


def preview_workbook(path: str | Path) -> ImportPreview:
    source = Path(path)
    book, sheet, kind = _open_sheet(source)
    try:
        headers, mapping, header_row = _headers(sheet, kind)
        count = sum(1 for index, row in enumerate(_rows(sheet, kind))
                    if index > header_row and any(v is not None and str(v).strip() for v in row))
        return ImportPreview(source, "Main Data", count, headers,
                             [c.label for c in COLUMNS if c.field not in mapping],
                             [h for h in headers if normalize_header(h) and normalize_header(h) not in {normalize_header(c.label) for c in COLUMNS}])
    finally:
        _close(book, kind)


def iter_import_rows(path: str | Path, warning: Callable[[int], None] | None = None) -> Iterator[dict[str, Any]]:
    source = Path(path)
    book, sheet, kind = _open_sheet(source)
    try:
        _, mapping, header_row = _headers(sheet, kind)
        for index, row in enumerate(_rows(sheet, kind)):
            if index <= header_row or not any(v is not None and str(v).strip() for v in row):
                continue
            result = {}
            for column in COLUMNS:
                value = row[mapping[column.field]] if column.field in mapping and mapping[column.field] < len(row) else None
                if column.kind == "number":
                    result[column.field] = _number(value, warning)
                elif column.kind == "date":
                    result[column.field] = _date(value, kind, book, warning)
                else:
                    result[column.field] = _text(value)
            yield result
    finally:
        _close(book, kind)


def _close(book, kind: str) -> None:
    if kind == "xlsx":
        book.close()
    else:
        book.release_resources()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if not text or text.casefold() in {"n/a", "na", "null"} else text


def _number(value: Any, warning: Callable[[int], None] | None) -> float | None:
    if value is None or (isinstance(value, str) and (not value.strip() or value.strip().casefold() in {"n/a", "na", "null"})):
        return None
    try:
        number = float(value.replace(",", "") if isinstance(value, str) else value)
        if not math.isfinite(number):
            raise ValueError
        return number
    except (TypeError, ValueError):
        if warning:
            warning(1)
        return None


def _date(value: Any, kind: str, book: Any, warning: Callable[[int], None] | None) -> str | None:
    if value is None or value == "":
        return None
    parsed: datetime | None = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            if kind == "xls":
                import xlrd
                parsed = xlrd.xldate.xldate_as_datetime(value, book.datemode)
            else:
                parsed = datetime(1899, 12, 30) + timedelta(days=float(value))
        except (ValueError, OverflowError):
            parsed = None
    else:
        text = str(value).strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
                try:
                    parsed = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
    if parsed is None:
        if warning:
            warning(1)
        return None
    return parsed.isoformat(sep=" ", timespec="seconds")
