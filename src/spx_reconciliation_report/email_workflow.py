"""Services for preparing selected Main Data records for email delivery."""
from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable

from .columns import COLUMNS, Column
from .email_service import EmailAccountRepository, EmailProviderService


class ReconciliationAttachment:
    MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def create(rows: list[dict[str, Any]], columns: Iterable[Column] | None = None,
               progress: Callable[[str, int, int], None] | None = None) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        selected_columns = tuple(COLUMNS if columns is None else columns)
        if not selected_columns:
            raise ValueError("At least one visible Main Data column is required to create the email attachment.")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reconciliation"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(selected_columns))}{len(rows) + 1}"
        sheet.append([column.label for column in selected_columns])
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="14283B")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        total = len(rows)
        for row_index, row in enumerate(rows, 1):
            values = []
            for column in selected_columns:
                value = row.get(column.field)
                if column.kind == "date" and value:
                    try:
                        value = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
                    except ValueError:
                        pass
                    if isinstance(value, datetime) and value.utcoffset() is not None:
                        value = value.isoformat(sep=" ")
                values.append(value)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
            if progress and (row_index % 250 == 0 or row_index == total):
                progress("writing", row_index, total)
        for column_index, column in enumerate(selected_columns, 1):
            letter = get_column_letter(column_index)
            max_length = max(
                (len(str(sheet.cell(row=row_index, column=column_index).value or ""))
                 for row_index in range(1, min(sheet.max_row, 201) + 1)), default=0
            )
            sheet.column_dimensions[letter].width = max(12, min(34, max_length + 3))
            if column.kind == "date":
                for cell in sheet[letter][1:]:
                    if cell.value:
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
        path = Path(tempfile.mkdtemp(prefix="spx-email-")) / f"SPX_Reconciliation_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
        if progress:
            progress("saving", total, total)
        workbook.save(path)
        workbook.close()
        return path


class EmailBatchService:
    def __init__(self, accounts: EmailAccountRepository, provider: EmailProviderService):
        self.accounts = accounts
        self.provider = provider

    def send(self, account_id: str, recipient_groups: list[dict[str, list[str]]], *,
             subject: str, html_body: str, attachment_path: Path,
             progress: Callable[[int, int], None] | None = None) -> list[dict[str, Any]]:
        results = []
        total = len(recipient_groups)
        attachment = attachment_path.read_bytes()
        for index, group in enumerate(recipient_groups, 1):
            result = {"to": group["to"], "cc": group["cc"], "bcc": group["bcc"],
                      "status": "SENT", "error": None}
            try:
                self.provider.send(
                    account_id, to=group["to"], cc=group["cc"], bcc=group["bcc"],
                    subject=subject, html_body=html_body,
                    attachment_name=attachment_path.name, attachment=attachment,
                )
            except Exception as exc:
                result["status"] = "FAILED"
                result["error"] = str(exc).splitlines()[0] or "Email could not be sent."
            results.append(result)
            if progress:
                progress(index, total)
        return results
